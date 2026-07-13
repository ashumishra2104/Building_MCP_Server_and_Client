"""
STAR Achievement Bank — parse, select, and validate stories.
No Streamlit imports. No side effects on import.
"""
import io
import json
import re
import os

import openpyxl

from src.helper import client, extract_text_from_pdf


# ── Parsing ───────────────────────────────────────────────────────

def parse_star_bank_excel(source) -> list[dict]:
    """
    Parse the STAR Achievement Bank Excel file.
    source: absolute file path str OR raw .xlsx bytes.
    Returns list of dicts: company, project_name, situation, task, action, result.
    Returns [] on any error.
    """
    try:
        if isinstance(source, (bytes, bytearray)):
            wb = openpyxl.load_workbook(io.BytesIO(source))
        else:
            wb = openpyxl.load_workbook(source)

        try:
            ws = wb["STAR Achievement Bank"]
        except KeyError:
            ws = wb.active

        stories = []
        for row in ws.iter_rows(min_row=4, values_only=True):
            company      = row[1]
            project_name = row[2]
            # col 3 = Amazon LP — intentionally skipped
            situation    = row[4]
            task         = row[5]
            action       = row[6]
            result       = row[7]

            if not company:
                continue

            stories.append({
                "company":      str(company).strip(),
                "project_name": str(project_name).strip() if project_name else "",
                "situation":    str(situation).strip()    if situation    else "",
                "task":         str(task).strip()         if task         else "",
                "action":       str(action).strip()       if action       else "",
                "result":       str(result).strip()       if result       else "",
            })
        return stories
    except Exception as e:
        print(f"[star_bank] parse_star_bank_excel error: {e}")
        return []


def parse_star_bank_pdf(pdf_bytes: bytes) -> list[dict]:
    """
    Parse a STAR bank from PDF bytes using OCR + GPT extraction.
    Returns list of dicts with same schema as parse_star_bank_excel.
    Returns [] on failure.
    """
    try:
        bio = io.BytesIO(pdf_bytes)
        bio.name = "star_bank.pdf"
        raw_text = extract_text_from_pdf(bio)
        if not raw_text.strip():
            return []

        prompt = f"""Extract all STAR achievement stories from the text below.
Return a JSON object with key "stories" containing an array. Each element must have exactly
these keys: company, project_name, situation, task, action, result.
Populate each field from the relevant section. If a field is missing use "".
Return ONLY valid JSON.

TEXT:
{raw_text[:8000]}"""

        resp = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            temperature=0,
            max_tokens=3000,
        )
        data = json.loads(resp.choices[0].message.content)
        stories = data.get("stories", [])
        required = {"company", "project_name", "situation", "task", "action", "result"}
        return [
            {k: str(s.get(k, "")).strip() for k in required}
            for s in stories if isinstance(s, dict)
        ]
    except Exception as e:
        print(f"[star_bank] parse_star_bank_pdf error: {e}")
        return []


# ── Selection ─────────────────────────────────────────────────────

def select_star_stories(jd_text: str, star_bank: list[dict], top_n: int = 3) -> list[dict]:
    """
    Use gpt-4o-mini (JSON mode, T=0) to pick the top N most JD-relevant stories.
    Sends only compact summaries to GPT — NOT full story text.
    Returns a subset of star_bank (0 to top_n items). Returns [] on no match or error.
    """
    if not star_bank or not jd_text:
        return []

    summaries = "\n".join(
        f"[{i}] {s['company']} | {s['project_name']}\n"
        f"    Result: {s['result'][:220]}"
        for i, s in enumerate(star_bank)
    )

    prompt = f"""You are evaluating which achievement stories from a candidate's experience bank
are most relevant to a Job Description.

Return a JSON object with key "indices" — an array of the top {top_n} most relevant story
indices (0-based integers), ordered by relevance. Return fewer if fewer are clearly relevant.
Return an empty array if none are a strong match.

JOB DESCRIPTION (first 2500 chars):
{jd_text[:2500]}

CANDIDATE STORIES:
{summaries}

Return ONLY valid JSON, e.g. {{"indices": [2, 7, 0]}}"""

    try:
        resp = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            temperature=0,
        )
        data = json.loads(resp.choices[0].message.content)
        indices = [i for i in data.get("indices", []) if isinstance(i, int) and 0 <= i < len(star_bank)]
        return [star_bank[i] for i in indices[:top_n]]
    except Exception as e:
        print(f"[star_bank] select_star_stories error: {e}")
        return []


# ── Validation ────────────────────────────────────────────────────

_METRIC_RE = re.compile(
    r'(\d[\d,\.]*\s*(?:crore|lakh|cr|%|x|X)?'
    r'|\d+\s*[→\-]\s*\d+'
    r'|\d+\s*(?:warehouses?|clients?|days?|weeks?|months?|orders?)'
    r'|\d+\.\d+%?)',
    re.IGNORECASE,
)

def validate_metrics(selected_stories: list[dict], html_output: str) -> list[tuple]:
    """
    Check that key metrics from selected stories appear verbatim in the generated HTML.
    Returns list of (project_name, missing_metric) tuples. Empty list = all good.
    """
    missing = []
    # Normalise HTML for comparison: strip tags, collapse whitespace
    plain = re.sub(r'<[^>]+>', ' ', html_output)
    plain = re.sub(r'\s+', ' ', plain)

    for s in selected_stories:
        for metric in set(_METRIC_RE.findall(s["result"])):
            metric_clean = metric.strip()
            if metric_clean and len(metric_clean) > 1:
                # Normalise whitespace in metric too before checking
                if metric_clean not in plain and metric_clean.replace(" ", "") not in plain.replace(" ", ""):
                    missing.append((s["project_name"], metric_clean))
    return missing


# ── STAR block builders ───────────────────────────────────────────

def build_resume_star_block(selected_stories: list[dict]) -> str:
    """
    Build the system-prompt injection block for tailor_resume().
    Tells GPT to weave story metrics into existing Experience bullets.
    """
    if not selected_stories:
        return ""

    block = f"\n\n    STAR ACHIEVEMENT BANK — WEAVING INSTRUCTIONS\n"
    block += f"    The following {len(selected_stories)} verified achievement story(ies) must inform your rewrite.\n"

    for i, s in enumerate(selected_stories, 1):
        block += f"""
    [STORY {i} — {s['company']} | {s['project_name']}]
    Situation : {s['situation']}
    Task      : {s['task']}
    Action    : {s['action']}
    Result    : {s['result']}
    [END STORY {i}]
"""

    block += """
    MANDATORY WEAVING RULES:
    1. Find bullets in the Experience section whose role/company/context overlaps with a story above.
    2. Where overlap exists, incorporate the story's confirmed metrics/outcomes into that bullet
       as supporting evidence. Blend naturally — do not append as a separate new sentence.
    3. Use ONLY numbers and facts that appear verbatim in the story.
       "55 crore" stays "55 crore". "99.5%" stays "99.5%". "500 → 700 STRs/day" stays as-is.
    4. Company binding: a story belongs to its named company. Only use its metrics in bullets
       for that company's role in the Experience section. Never cross-contaminate companies.
    5. Do NOT create a new section or sub-heading for these stories. Weave into existing bullets only.
    6. If no existing bullet is a natural home for a story's metrics — skip that story entirely.
    7. All anti-fabrication rules above still apply on top of these instructions.
    """
    return block


def build_cover_letter_star_block(selected_stories: list[dict]) -> str:
    """
    Build the system-prompt injection block for generate_cover_letter().
    Tells GPT to use story metrics for the HIGHLIGHT_1–4 callout boxes.
    """
    if not selected_stories:
        return ""

    block = f"\n\n    STAR ACHIEVEMENT BANK — HIGHLIGHT BOX INSTRUCTIONS\n"
    block += f"    The following {len(selected_stories)} verified achievement story(ies) contain real, confirmed metrics.\n"

    for i, s in enumerate(selected_stories, 1):
        block += f"""
    [STORY {i} — {s['company']} | {s['project_name']}]
    Result: {s['result']}
    [END STORY {i}]
"""

    block += """
    MANDATORY RULES FOR HIGHLIGHT BOXES:
    1. For {{HIGHLIGHT_1_LABEL}} through {{HIGHLIGHT_4_LABEL}} and their matching DETAIL fields,
       prioritise metrics from the STAR stories above over generic resume achievements.
    2. Use ONLY numbers and outcomes that appear verbatim in the stories. Never paraphrase metrics.
    3. Company binding: only use a story's metrics if the story's company matches a role in the resume.
    4. If fewer than 4 STAR metrics are clearly applicable, fill remaining highlight boxes
       from the resume as normal — do not force unrelated story metrics.
    5. All anti-fabrication rules above still apply.
    """
    return block
